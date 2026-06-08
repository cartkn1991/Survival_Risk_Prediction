#!/usr/bin/env python3
"""Write methods summary document (Word .docx) for GWAS pipeline."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def _add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _load_json(path: Path) -> dict:
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return {}


def build_doc(out_path: Path) -> None:
    base = Path("feature_importance/gwas")
    rep_dir = base / "replicated"

    disc = _load_json(base / "discovery_summary.json")
    mort = _load_json(base / "mortality_fhs_summary.json")
    cross = _load_json(base / "cross_compare_summary.json")
    cox_rep = _load_json(base / "snp_cox_replication_summary.json")
    follow = _load_json(rep_dir / "replicated_snp_followup_summary.json")
    minigwas_path = Path("feature_importance/minigwas/summary.json")
    mini = _load_json(minigwas_path)

    doc = Document()
    title = doc.add_heading("AESurv GWAS and mortality genetics — methods summary", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_para(doc, f"Generated: {date.today().isoformat()}")
    _add_para(doc, "Cohorts: Framingham Heart Study (FHS, discovery n≈3,209) and Women's Health Initiative (WHI, replication n≈510).")

    _add_heading(doc, "1. Overview", 1)
    _add_para(
        doc,
        "We evaluated genetic associations with AESurv predicted mortality risk (log_h) and with observed "
        "mortality in a staged design: (i) mini-GWAS on model-selected SNPs in outlier groups; "
        "(ii) genome-wide discovery in FHS; (iii) WHI replication; (iv) cross-phenotype comparison of log_h vs mortality GWAS; "
        "(v) per-SNP Cox proportional hazards in FHS and WHI; (vi) annotation and outlier stratification for replicated loci.",
    )

    _add_heading(doc, "2. Data and genotypes", 1)
    _add_bullets(doc, [
        "Genotypes: additive dosage (0/1/2) from NPZ bundles (~1.04M SNPs per cohort).",
        "FHS bundle: vae_cox_cache/bundles/FHS_*_cpgall_snpall.npz (n=3,209).",
        "WHI bundle: vae_cox_cache/bundles/WHI_*_cpgall_snpall.npz (n=510).",
        "AESurv log_h: runs/aesurv_aux_grid/age12.00_cell1.00/aesurv_aux_risk.npz (FHS train/val split reproduced; WHI test).",
        "Mortality: time and event from bundle NPZ; chronological age and sex from cohort parquet metadata.",
    ])

    _add_heading(doc, "3. Mini-GWAS (outlier contrast, restricted SNP set)", 1)
    _add_para(doc, "Script: minigwas_outliers.py. Output: feature_importance/minigwas/.")
    _add_bullets(doc, [
        f"Input SNPs: {mini.get('n_significant_input_snps', '11,939')} AESurv risk-gradient significant SNPs (significant_risk.csv).",
        "Groups (pooled FHS+WHI): resilient = bottom 10% of age+cohort-adjusted log_h residual; accelerated = top 10%.",
        "Model: logistic regression per SNP — logit(P(accelerated)) ~ SNP dose + cohort.",
        f"Samples: {mini.get('n_resilient', 372)} resilient, {mini.get('n_accelerated', 372)} accelerated; {mini.get('n_after_qc', '—')} SNPs after QC.",
        f"Genomic inflation: lambda_GC = {mini.get('lambda_gc', '—')}.",
        f"Bonferroni hits: {mini.get('n_pass_bonferroni', 0)}; BH FDR<5%: {mini.get('n_pass_bh_fdr_5pc', 0)}.",
        "Interpretation: exploratory extreme-group scan; not primary discovery due to low power and multiple testing.",
    ])

    _add_heading(doc, "4. Full GWAS — age-adjusted log_h residual (FHS discovery)", 1)
    _add_para(doc, "Script: gwas/run_full_gwas_pipeline.py --stage discovery. Output: discovery_fhs_results.csv.")
    _add_bullets(doc, [
        "Phenotype: residual from OLS log_h ~ age + sex within FHS.",
        f"SNPs tested: {disc.get('n_snps_tested', '—'):,} (post MAF≥1%, missing≤5%).",
        f"lambda_GC = {disc.get('lambda_gc', '—')}.",
        f"Suggestive hits (p≤1×10⁻⁵): {disc.get('n_suggestive', '—')}; FDR q≤0.05: {disc.get('n_fdr_5pc', 0)}.",
        "WHI replication: clumped FHS leads tested with same residual model; fixed-effect meta (replication_results.csv).",
        "Outlier stratification: mean dosage across resilient/middle/accelerated on FHS (post hoc).",
        f"Enrichment vs 11k gradient SNPs: overlap at suggestive threshold per enrichment_summary.json.",
    ])

    _add_heading(doc, "5. Full GWAS — mortality (FHS)", 1)
    _add_para(doc, "Script: gwas/run_fhs_mortality_gwas.py. Output: mortality_fhs_results.csv.")
    _add_bullets(doc, [
        "Phenotype: binary death event; linear model event ~ SNP + age_z + sex (screening surrogate for Cox).",
        f"Events: {mort.get('n_events', 575)} / {mort.get('n_samples', 3209)}.",
        f"SNPs tested: {mort.get('n_snps_tested', '—'):,}; lambda_GC = {mort.get('lambda_gc', '—')}.",
        f"Genome-wide significant (p<5×10⁻⁸): {mort.get('n_genomewide', 0)}.",
    ])

    _add_heading(doc, "6. Cross-phenotype comparison (log_h vs mortality)", 1)
    _add_para(doc, "Script: gwas/cross_compare_gwas_hits.py.")
    _add_bullets(doc, [
        f"log_h clumped leads: {cross.get('n_logh_leads', '—')}; mortality leads: {cross.get('n_mortality_leads', '—')}.",
        f"Exact SNP overlap: {cross.get('n_exact_overlap', 0)}; LD-window pairs ({cross.get('ld_window_bp', 250000)//1000} kb): {cross.get('n_window_pairs', 0)}.",
        "Union table: cross_compare_union.csv with betas/p for both phenotypes per SNP.",
    ])

    _add_heading(doc, "7. Per-SNP Cox proportional hazards", 1)
    _add_para(doc, "Script: gwas/run_fhs_snp_cox.py (--cohort FHS or WHI).")
    _add_bullets(doc, [
        "Model: Surv(time, event) ~ SNP + age [+ sex if variable]; lifelines CoxPHFitter, penalizer=0.01.",
        "FHS: 20 union SNPs tested; 7 significant at p<0.05 (fhs_snp_cox_results.csv).",
        "WHI replication: 7 FHS-significant SNPs; 6 genotyped in WHI (4:183579487_T missing).",
        f"Strict replication (both p<0.05, same sign): {cox_rep.get('n_replicated_strict', 2)}.",
        "Merged table: snp_cox_replication.csv with meta-analysis coefficients.",
    ])

    _add_heading(doc, "8. Replicated SNPs — annotation and outlier stratification", 1)
    _add_para(doc, "Script: gwas/replicated_snp_followup.py. Output: feature_importance/gwas/replicated/.")
    replicated = [r for r in cox_rep.get("rows", []) if r.get("replicated")]
    if replicated:
        for r in replicated:
            _add_para(
                doc,
                f"• {r['snp']}: FHS HR={r['fhs_hr_per_allele']:.2f} (p={r['fhs_p']:.2e}), "
                f"WHI HR={r['whi_hr_per_allele']:.2f} (p={r['whi_p']:.3f}), meta HR={r['meta_hr']:.2f} (p={r['meta_p']:.3f}).",
            )
    if follow.get("annotation"):
        _add_para(doc, "Annotation (Ensembl genes ±5 kb GWAS catalog rsIDs; AESurv 11k panel membership):", bold=True)
        for a in follow["annotation"]:
            _add_para(
                doc,
                f"  {a['snp']}: genes={a.get('gene_ensembl', '') or '—'}; "
                f"nearby GWAS rsIDs={a.get('gwas_rsids_near', '') or '—'}; "
                f"in gradient 11k={a.get('in_aesurv_gradient_11k', False)}.",
            )
    if follow.get("stratification"):
        _add_para(doc, "FHS dosage trend (resilient → accelerated):", bold=True)
        for s in follow["stratification"]:
            _add_para(
                doc,
                f"  {s['snp']}: resilient={s['mean_dose_resilient']:.3f}, middle={s['mean_dose_middle']:.3f}, "
                f"accelerated={s['mean_dose_accelerated']:.3f}; trend r={s['trend_r']:.3f}, p={s['trend_p']:.2e}; "
                f"expected direction={s['expected_direction_in_tails']}.",
            )

    _add_heading(doc, "8.1 Results — replicated SNPs (interpretation)", 2)
    _add_para(
        doc,
        "11:127103604_C (chr11:127103604): This locus shows the strongest integrated evidence. In FHS Cox, "
        "each additional allele was associated with 15% higher mortality (HR=1.15, p=0.034); in WHI the estimate "
        "was larger (HR=1.76, p=0.046), and the fixed-effect meta-analysis yielded HR=1.18 (meta p=0.012). "
        "The variant is not among the 11,939 AESurv gradient-significant SNPs, indicating that it tags "
        "population mortality risk rather than in-model feature importance. Ensembl overlap did not return a "
        "protein-coding symbol at the exact position; a nearby GWAS Catalog entry is rs11220776 (±5 kb). "
        "On FHS, mean allele dosage increased monotonically from resilient (0.40) to middle (0.51) to accelerated "
        "(0.54) outlier groups (Pearson trend r=0.058, p=0.001), consistent with the direction of the Cox hazard "
        "ratio. This SNP is the best candidate mortality biomarker from the current pipeline: replicated across "
        "cohorts and aligned with age-adjusted high-risk tails.",
    )
    _add_para(
        doc,
        "12:46600556_T (chr12:46600556): This locus also replicated in two-cohort Cox models (FHS HR=1.62, p=1.3×10⁻⁴; "
        "WHI HR=1.27, p=0.032; meta HR=1.41, meta p=3.4×10⁻⁵). It maps near SLC38A4-AS1 (Ensembl) with a nearby "
        "catalog SNP rs4768737 and is likewise absent from the 11k gradient panel. However, FHS outlier "
        "stratification did not show the expected dosage gradient across resilient, middle, and accelerated groups "
        "(means 0.065, 0.076, 0.052; trend p=0.51). Thus 12:46600556_T is a statistically robust replicated mortality "
        "association but weaker support for enrichment in AESurv-defined accelerated aging tails on FHS. It may reflect "
        "a direct mortality effect partly decoupled from age-adjusted predicted log_h, or limited power/rare-variant "
        "behavior at this locus (MAF≈3.6% in FHS).",
    )
    _add_para(
        doc,
        "Summary contrast: Both replicated SNPs are independent of the AESurv gradient feature set. "
        "11:127103604_C combines replication, meta-analysis significance, and outlier-tail concordance; "
        "12:46600556_T combines replication with annotation near SLC38A4-AS1 but without clear tail enrichment. "
        "Neither locus was genome-wide significant in the linear mortality screen (p~10⁻⁶); prioritization rested "
        "on strict per-cohort Cox replication criteria.",
    )

    _add_heading(doc, "9. Key files (feature_importance/gwas/)", 1)
    files = [
        "discovery_fhs_results.csv, discovery_fhs_manhattan.png, discovery_summary.json",
        "mortality_fhs_results.csv, mortality_fhs_summary.json",
        "cross_compare_union.csv, cross_compare_summary.json",
        "fhs_snp_cox_results.csv, whi_snp_cox_results.csv, snp_cox_replication.csv",
        "replicated/replicated_snp_annotation.csv, replicated_snp_stratification.csv",
        "minigwas/: minigwas_results.csv, summary.json (separate folder)",
    ]
    _add_bullets(doc, files)

    _add_heading(doc, "10. Conclusions", 1)
    _add_bullets(doc, [
        "Mini-GWAS and log_h GWAS interrogate model-related risk; mortality GWAS and Cox test observed death.",
        "No genome-wide significant loci at p<5×10⁻⁸; several suggestive loci (p~10⁻⁶) in FHS.",
        "Two SNPs replicate in FHS and WHI Cox: 12:46600556_T and 11:127103604_C (same direction, both p<0.05).",
        "11:127103604_C is the priority mortality biomarker candidate (meta p≈0.012; outlier dosage trend p≈0.001).",
        "12:46600556_T shows strong Cox replication (meta p≈3.4×10⁻⁵) but flat outlier-tail dosage on FHS.",
        "log_h GWAS hits largely do not overlap mortality GWAS or the 11k gradient SNP panel.",
        "In-sample log_h training on FHS limits causal claims; mortality Cox + WHI replication is the stronger genetics arm.",
    ])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Wrote {out_path}")


def main() -> None:
    import argparse
    p = argparse.ArgumentParser()
    p.add_argument("--out", default="feature_importance/gwas/AESurv_GWAS_methods_summary.docx")
    args = p.parse_args()
    build_doc(Path(args.out))


if __name__ == "__main__":
    main()
